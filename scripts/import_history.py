"""One-time historical import (spec Appendix). Not part of the bot; run once, then discard.

    python -m scripts.import_history --dry-run     # parse + show what would be loaded, no LLM, no writes
    python -m scripts.import_history               # load into the database, tag muscle groups, replay notes

Owner-confirmed corrections to the spreadsheet (2026-10-04):
- Strength "Порядок сессии" runs NEWEST -> oldest (1 = most recent). Cardio/body sheets run oldest -> newest.
- Chest session 4 "Pullover Machine" 20 kg is a typo for "Reverse Pec Deck".
- Chest session 3 Pullover Machine was a one-off substitution (Reverse Pec Deck broken) -> archived.
- Knee issue existed and is resolved (legs session 3 mentions technique fixing the knee).
- Treadmill baseline: гора 8%/5.5, предгорье 4%/6.5.
"""

import argparse
import asyncio
import re
import shutil
import sys
from collections import defaultdict
from datetime import time as dtime
from datetime import timedelta
from pathlib import Path

import openpyxl
from sqlalchemy import select

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.config import now  # noqa: E402
from app.context import fmt_date, fmt_num, fmt_perf, time_block  # noqa: E402
from app.db import db_session, init_db  # noqa: E402
from app.models import (BodyMetric, CardioRecord, Exercise, Injury, Performance, Profile, Split,  # noqa: E402
                        TrainingSession)
from app.synthesis import (apply_session_result, cardio_text, profile_snapshot, synthesize_exercise_note,  # noqa: E402
                           synthesize_session)
from app.tools import classify_muscle_groups  # noqa: E402

# Your filled-in copy of training_history_import_template.xlsx (kept out of git: it holds personal data)
XLSX = Path(__file__).resolve().parent.parent / "training_history_import_template_2.xlsx"

# Placeholder dating (spec Appendix "Dates"): newest session of each split lands in the past week,
# older sessions of the same split step back one week each.
SPLIT_DAYS_AGO = {"Спина": 6, "Грудь": 4, "Ноги": 2}
TREADMILL_BASELINE = [
    {"label": "гора", "incline_pct": 8, "speed_kmh": 5.5},
    {"label": "предгорье", "incline_pct": 4, "speed_kmh": 6.5},
]
ARCHIVED_ONE_OFFS = {"Pullover Machine"}


def seconds(value) -> int:
    if value is None:
        return 0
    if isinstance(value, dtime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    parts = [int(p) for p in str(value).split(":")]
    while len(parts) < 3:
        parts.insert(0, 0)
    return parts[0] * 3600 + parts[1] * 60 + parts[2]


def parse_profile(text: str | None) -> list[dict]:
    """'гора 8%/5.5, предгорье 4%/6.7' -> [{label, incline_pct, speed_kmh}, ...]"""
    out = []
    for chunk in (text or "").split(","):
        m = re.search(r"([^\d,]*?)\s*(\d+(?:[.,]\d+)?)\s*%\s*/\s*(\d+(?:[.,]\d+)?)", chunk)
        if m:
            out.append({"label": m.group(1).strip(), "incline_pct": float(m.group(2).replace(",", ".")),
                        "speed_kmh": float(m.group(3).replace(",", "."))})
    return out


def read_sheets():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    strength, cardio, body = [], [], []
    for row in wb["История силовых"].iter_rows(min_row=2, values_only=True):
        split, order, name, weight, reps, sets, comment = (list(row) + [None] * 7)[:7]
        if not split or not isinstance(order, (int, float)) or not name:
            continue
        name = str(name).strip()
        if split == "Грудь" and int(order) == 4 and name == "Pullover Machine":
            name = "Reverse Pec Deck"  # owner-confirmed typo
        strength.append({"split": split.strip(), "order": int(order), "name": name,
                         "weight": float(weight) if weight is not None else None,
                         "reps": int(reps) if reps is not None else None,
                         "sets": int(sets) if sets is not None else None,
                         "comment": (comment or "").strip()})
    for row in wb["История кардио"].iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        order, modality, duration, profile, level, z1, z2, z3, z4, z5, kcal, carbs, protein, fat, comment = row[:15]
        modality = "elliptical" if str(modality).strip().lower().startswith("орб") else "treadmill"
        cardio.append({"order": int(order), "modality": modality, "duration": duration,
                       "intervals": parse_profile(profile) if modality == "treadmill" else [],
                       "level": level if modality == "elliptical" else None,
                       "zones": [seconds(z) for z in (z1, z2, z3, z4, z5)], "kcal": kcal,
                       "carbs": carbs, "protein": protein, "fat": fat, "comment": (comment or "").strip()})
    for row in wb["Замеры тела"].iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], (int, float)):
            body.append({"order": int(row[0]), "weight": row[1], "muscle": row[2], "fat": row[3]})
    return strength, sorted(cardio, key=lambda c: c["order"]), sorted(body, key=lambda b: b["order"])


def plan_sessions(strength):
    """Groups rows into dated sessions, oldest first."""
    today = now().replace(hour=19, minute=0, second=0, microsecond=0)
    sessions = defaultdict(list)
    for r in strength:
        sessions[(r["split"], r["order"])].append(r)
    dated = []
    for (split, order), rows in sessions.items():
        days_ago = SPLIT_DAYS_AGO.get(split, 7) + 7 * (order - 1)  # order 1 = newest
        dated.append({"split": split, "order": order, "date": today - timedelta(days=days_ago), "rows": rows})
    return sorted(dated, key=lambda s: s["date"])


async def main(dry_run: bool) -> None:
    strength, cardio, body = read_sheets()
    sessions = plan_sessions(strength)
    # Cardio (oldest -> newest) is attached to the most recent workouts, in date order.
    recent = sessions[-len(cardio):] if cardio else []
    for sess, c in zip(recent, cardio):
        sess["cardio"] = c
    names = sorted({r["name"] for r in strength})

    print(f"{len(strength)} strength rows, {len(sessions)} sessions, {len(names)} unique exercises, "
          f"{len(cardio)} cardio, {len(body)} body metrics\n")
    for s in sessions:
        c = s.get("cardio")
        print(f"{fmt_date(s['date'])}  {s['split']:6} (sheet order {s['order']}): "
              + ", ".join(f"{r['name'].split(' (')[0]} {fmt_num(r['weight'])}×{r['reps']}×{r['sets']}" for r in s["rows"])
              + (f"  + cardio {c['modality']}" if c else ""))
    if dry_run:
        print("\nTreadmill intervals parsed:", [c["intervals"] for c in cardio])
        return

    init_db()
    with db_session() as db:
        if db.scalar(select(Performance.id).limit(1)):
            raise SystemExit("Database already has performance history - import refused (it's a one-time load).")

    print("\nTagging muscle groups (one call per unique exercise)...")
    groups = dict(zip(names, await asyncio.gather(*(classify_muscle_groups(n) for n in names))))
    for n in names:
        print(f"  {n}: {groups[n]}")

    with db_session() as db:
        p = db.get(Profile, 1)
        p.treadmill_baseline = TREADMILL_BASELINE
        p.default_cardio_modality = "treadmill"
        latest_per_split = {}
        for s in sessions:
            latest_per_split[s["split"]] = {r["name"] for r in s["rows"]}
        for n in names:
            active = n not in ARCHIVED_ONE_OFFS and any(n in v for v in latest_per_split.values())
            db.add(Exercise(name=n, muscle_groups=groups[n], status="active" if active else "archived",
                            created_at=sessions[0]["date"]))
        legs = [s for s in sessions if s["split"] == "Ноги"]
        knee_date = legs[0]["date"] if legs else sessions[0]["date"]
        db.add(Injury(description="Боль в колене при жиме ногами (ушла после исправления техники)", body_area="колено",
                      muscle_groups=["quads"], status="resolved", created_at=knee_date - timedelta(days=7),
                      last_mentioned_at=knee_date, resolved_at=knee_date, temporary_substitutions=[], reverts_proposed=True))
        for b in body:
            db.add(BodyMetric(measured_at=sessions[-1]["date"] + timedelta(hours=-10), weight_kg=b["weight"],
                              muscle_kg=b["muscle"], fat_pct=b["fat"]))

    print("\nReplaying synthesis session by session (oldest -> newest)...")
    for s in sessions:
        with db_session() as db:
            split = db.scalar(select(Split).where(Split.name == s["split"]))
            ts = TrainingSession(started_at=s["date"], last_activity_at=s["date"], closed_at=s["date"], status="closed",
                                 primary_split_id=split.id, transcript=[], applied_changes=[], imported=True)
            db.add(ts)
            db.flush()
            perfs = []
            for r in s["rows"]:
                ex = db.scalar(select(Exercise).where(Exercise.name == r["name"]))
                perf = Performance(session_id=ts.id, exercise_id=ex.id, performed_at=s["date"], weight_kg=r["weight"],
                                   reps=r["reps"], sets=r["sets"], comment=r["comment"])
                db.add(perf)
                perfs.append((ex, perf))
            c = s.get("cardio")
            cardio_row = None
            if c:
                cardio_row = CardioRecord(session_id=ts.id, performed_at=s["date"] + timedelta(hours=1), modality=c["modality"],
                                          duration_min=c["duration"], zone_seconds=c["zones"], calories=c["kcal"],
                                          carbs_pct=c["carbs"], protein_pct=c["protein"], fat_pct=c["fat"],
                                          treadmill_intervals=c["intervals"], elliptical_level=c["level"], comment=c["comment"])
                db.add(cardio_row)
            db.flush()

            context = "\n".join(f"- {ex.name}: {fmt_perf(perf)}" for ex, perf in perfs)
            notes = await asyncio.gather(*(
                synthesize_exercise_note(exercise_name=ex.name, muscle_groups=ex.muscle_groups, current_note=ex.rolling_note,
                                         session_rows=[fmt_perf(perf)], date_text=fmt_date(s["date"]),
                                         session_context=f"Imported historical session (no conversation). All results that day:\n{context}")
                for ex, perf in perfs))
            for (ex, _), note in zip(perfs, notes):
                ex.rolling_note = note

            session_text = f"Date: {fmt_date(s['date'])}; primary split: {split.name}\n{context}"
            if cardio_row is not None:
                session_text += "\n" + cardio_text(cardio_row)
            result = await synthesize_session(
                profile_text=profile_snapshot(db), session_text=session_text, already_applied=[],
                split_text=f"## Primary split: {split.name}\nCurrent split note: {split.rolling_note or '—'}")
            result["goals"], result["injuries"], result["body_metrics"] = None, [], []  # import seeds notes only
            apply_session_result(db, result, split, s["date"])
            print(f"  {fmt_date(s['date'])} {split.name}: {len(perfs)} notes updated" + (" + cardio" if c else ""))

    from app.db import engine
    engine.dispose()
    db_path = engine.url.database
    if db_path:
        snapshot = Path(__file__).resolve().parent.parent / "data" / "after_import.db"
        shutil.copy(db_path, snapshot)
        print(f"\nSnapshot for model tests saved: {snapshot}")

    print("\nDone. Exercise statuses (review these!):")
    with db_session() as db:
        for ex in db.scalars(select(Exercise).order_by(Exercise.status, Exercise.name)).all():
            print(f"  [{ex.status}] {ex.name}\n      {ex.rolling_note}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    asyncio.run(main(parser.parse_args().dry_run))
